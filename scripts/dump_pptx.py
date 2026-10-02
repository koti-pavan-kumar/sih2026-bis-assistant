"""Dump every text run in a pptx, recursing groups/tables, with slide+shape indices.

Usage: python scripts/dump_pptx.py <path-to.pptx> [--grep PATTERN]
"""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def iter_runs(shapes, prefix=""):
    for i, sh in enumerate(shapes):
        path = f"{prefix}{i}"
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_runs(sh.shapes, prefix=f"{path}.")
        elif sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    for p, para in enumerate(cell.text_frame.paragraphs):
                        for run in para.runs:
                            if run.text.strip():
                                yield f"{path}[{r},{c}]", run.text
        elif sh.has_text_frame:
            for p, para in enumerate(sh.text_frame.paragraphs):
                for run in para.runs:
                    if run.text.strip():
                        yield f"{path}p{p}", run.text


def main():
    path = sys.argv[1]
    grep = None
    if "--grep" in sys.argv:
        grep = sys.argv[sys.argv.index("--grep") + 1]
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    prs = Presentation(path)
    for si, slide in enumerate(prs.slides, 1):
        hits = []
        for shp, text in iter_runs(slide.shapes):
            if grep is None or grep.lower() in text.lower():
                hits.append((shp, text))
        if hits:
            print(f"=== Slide {si} ===")
            for shp, text in hits:
                print(f"  [{shp}] {text}")


if __name__ == "__main__":
    main()
