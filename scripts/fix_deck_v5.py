"""Patch ManakMitra_SIH2026_v5.pptx credibility issues.

Replaces absolute/unsupportable claims with measured, defensible statements,
fixes the dead demo URL, and corrects the budget line. Run-aware replacement
handles strings that span multiple runs. Verifies every replacement landed.

Usage: python scripts/fix_deck_v5.py [--dry-run]
"""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DECK = "ManakMitra_SIH2026_v5.pptx"
LIVE_URL = "sih2026-bis-assistant-ebon.vercel.app"

REPLACEMENTS = [
    # 2. absolute hallucination claim -> measured
    (" - ONNX + FAISS, hallucination-free, cited responses",
     " - ONNX + FAISS, citation-verified, source-grounded responses"),
    # 3. absolute claim in key features
    ("Hallucination-Free RAG", "Citation-Verified RAG"),
    # 5. Trust bullet -> measured numbers from scripts/eval_rag.py
    ("Source-grounded RAG ensures responses cite official BIS documents. No hallucinations.",
     "Cited sources on every answer; 97.8% of IS references verified against retrieved text (28-question eval harness in repo)."),
    # 6. benefits bullet -> measured
    ("All responses cite official BIS documents. Zero hallucinations.",
     "All responses cite official BIS documents; 97.8% of references verified in a 28-question eval."),
    # 5. market demand: "no free AI tool" absolute -> defensible comparison
    ("63M+ MSMEs in India struggle with BIS compliance. No free AI tool exists for this.",
     "63M+ MSMEs in India struggle with BIS compliance. Free AI tools answer generically - none indexes official IS text with citations."),
    # 5. partnerships presented as fact -> honest outreach status
    ("Potential BIS official data provider status. Revenue share on certified leads.",
     "No committed partners yet - outreach planned to BIS-recognized labs and MSME bodies; freemium subscriptions as the revenue model."),
    # 3 & 7. dead demo URL (404) -> live deployment
    ("manakmitra-frontend.vercel.app", LIVE_URL),
    # 6. BIS standards are public - budget line said "BIS Data Licensing"
    ("BIS Data Licensing", "Data Engineering"),
    # 4. mock answer cited IS 8112 which is not in the indexed corpus
    ("IS 269, IS 455, IS 8112", "IS 269, IS 455, IS 1489"),
]


def iter_paragraphs(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_paragraphs(sh.shapes)
        elif sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs
        elif sh.has_text_frame:
            yield from sh.text_frame.paragraphs


def replace_in_paragraph(para, old, new):
    """Replace `old` with `new` even when it spans several runs."""
    runs = list(para.runs)
    texts = [r.text for r in runs]
    full = "".join(texts)
    i = full.find(old)
    if i == -1:
        return False
    j = i + len(old)
    pos = 0
    placed = False
    for r, t in zip(runs, texts):
        s, e = pos, pos + len(t)
        pos = e
        if e <= i or s >= j:
            continue  # run untouched
        pre = t[: max(0, i - s)]
        suf = t[max(0, min(len(t), j - s)):]
        if not placed:
            r.text = pre + new + suf
            placed = True
        else:
            r.text = pre + suf
    return True


def main():
    dry = "--dry-run" in sys.argv
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    prs = Presentation(DECK)
    counts = {old: 0 for old, _ in REPLACEMENTS}
    for slide in prs.slides:
        for para in iter_paragraphs(slide.shapes):
            for old, new in REPLACEMENTS:
                while replace_in_paragraph(para, old, new):
                    counts[old] += 1
    for old, _ in REPLACEMENTS:
        n = counts[old]
        print(f"{'DRY ' if dry else ''}{'OK ' if n else 'MISS'} {n}x  {old[:70]}")
    if not dry:
        if all(counts.values()):
            prs.save(DECK)
            print(f"\nSaved {DECK}")
        else:
            print("\nNOT saving - some replacements missed")


if __name__ == "__main__":
    main()
