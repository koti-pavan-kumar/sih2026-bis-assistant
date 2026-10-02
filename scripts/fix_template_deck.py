"""Text fixes inside the submitted template deck 'SIH format My Team.pptx'.

1. Slide 2: stale "Telugu, English and Hindi" multilingual claim -> 22 languages.
2. Slide 6 (Research & References): append a prior-art/regulatory research
   entry (the slide previously listed only tools/vendors).

Backup: 'SIH format My Team.backup.pptx' must exist (made before image refresh).
Usage: python scripts/fix_template_deck.py [--dry-run]
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).parent.parent
WORKSPACE = ROOT.parent
DECK = WORKSPACE / "SIH format My Team.pptx"

OLD_LANG = " in Telugu, English and Hindi, making BIS information easier to access and understand."
NEW_LANG = " in 22 Indian languages, making BIS information easier to access and understand."

RESEARCH_NAME = "Prior Art & Regulatory Review:"
RESEARCH_DESC = ("Reviewed BIS Act 2016 and BIS conformity assessment rules for the legal basis, "
                 "and compared general AI assistants - none indexes IS text with section-level citations.")


def replace_in_paragraph(para, old, new):
    runs = list(para.runs)
    texts = [r.text for r in runs]
    full = "".join(texts)
    i = full.find(old)
    if i == -1:
        return False
    j = i + len(old)
    pos = 0
    placed = False
    for r, t in zip(runs, texts):
        s, e = pos, pos + len(t)
        pos = e
        if e <= i or s >= j:
            continue
        pre = t[: max(0, i - s)]
        suf = t[max(0, min(len(t), j - s)):]
        if not placed:
            r.text = pre + new + suf
            placed = True
        else:
            r.text = pre + suf
    return True


def main():
    dry = "--dry-run" in sys.argv
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    prs = Presentation(str(DECK))

    # 1. stale language claim (slide 2)
    fixed_lang = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                if replace_in_paragraph(para, OLD_LANG, NEW_LANG):
                    fixed_lang += 1
    print(f"{'DRY ' if dry else ''}language claim fixes: {fixed_lang}")

    # 2. research entry on the references slide (slide 6)
    appended = 0
    s6 = list(prs.slides)[5]
    for sh in s6.shapes:
        if not (sh.has_text_frame and "BIS Official Website:" in sh.text_frame.text):
            continue
        tf = sh.text_frame
        if any(RESEARCH_NAME in p.text for p in tf.paragraphs):
            print("research entry already present")
            continue
        paras = tf.paragraphs
        name_proto = next(p._p for p in paras if p.text.strip().endswith("Google Maps Platform:"))
        desc_proto = next(p._p for p in paras if "testing laboratory locations" in p.text)
        if dry:
            print("DRY would append:", RESEARCH_NAME, "/", RESEARCH_DESC[:40])
            continue
        tf._txBody.append(copy.deepcopy(name_proto))
        tf._txBody.append(copy.deepcopy(desc_proto))
        new_paras = tf.paragraphs[-2:]
        new_paras[0].runs[0].text = RESEARCH_NAME
        desc_runs = new_paras[1].runs
        if len(desc_runs) >= 2:
            desc_runs[-1].text = RESEARCH_DESC
        else:
            desc_runs[0].text = RESEARCH_DESC
        appended += 1
    print(f"{'DRY ' if dry else ''}research entries appended: {appended}")

    if not dry and fixed_lang and appended:
        prs.save(str(DECK))
        print(f"Saved {DECK.name}")
    elif not dry and fixed_lang:
        prs.save(str(DECK))
        print(f"Saved {DECK.name} (language only)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
